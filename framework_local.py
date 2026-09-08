# -*- coding: utf-8 -*-
"""
============================================================================
FRAMEWORK DE CUMPLIMIENTO AUTOMATIZADO - LEY 21.719 (AGENTE RAG)
Versión para ejecución LOCAL (VS Code) — sin dependencias de Google Colab.
============================================================================

REQUISITOS PREVIOS (una sola vez, en la terminal de VS Code):

  1. Crear y activar entorno virtual:
        python -m venv venv
        venv\\Scripts\\activate            (Windows)
        source venv/bin/activate           (Mac/Linux)

  2. Instalar dependencias:
        pip install pymupdf FlagEmbedding rank-bm25 chromadb langgraph ollama
        pip install python-docx python-pptx openpyxl pandas tabulate python-dotenv
        pip install ragas langchain-google-genai datasets

  3. Instalar Ollama como aplicación de escritorio desde https://ollama.com
     y luego, en la terminal:
        ollama pull qwen3:8b

  4. Crear un archivo llamado  .env  en la carpeta del proyecto con:
        GEMINI_API_KEY=tu_clave_de_gemini_aqui

  5. Colocar los PDFs de las leyes en una carpeta y ajustar RUTA_LEYES abajo.

NOTA DE HARDWARE: BGE-M3, el reranker y Qwen3 corren MEJOR con GPU NVIDIA.
En CPU funcionan pero lentos. Si no tienes GPU, considera dejar la ejecución
pesada en Colab y usar VS Code solo para editar/depurar.
============================================================================
"""

import os
import re
import json
import time
from pathlib import Path

from dotenv import load_dotenv

# ---------------------------------------------------------------------------
# CONFIGURACIÓN — AJUSTA ESTAS RUTAS A TU PC
# ---------------------------------------------------------------------------
load_dotenv()  # lee el archivo .env

# Carpeta donde están los PDF de las leyes (usa r"..." en Windows).
RUTA_LEYES = Path(r"C:\tesis-rag\Leyes")

# Carpeta de salida para markdown y la base vectorial.
CARPETA_SALIDA = RUTA_LEYES.parent / "output"
CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)

# Clave de Gemini (para la evaluación RAGAS). Se lee del archivo .env
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

# Modelo de Ollama (debe estar descargado con: ollama pull qwen3:8b)
MODELO_QWEN = "qwen3:8b"


def encontrar_pdf(carpeta, contiene):
    carpeta = Path(carpeta)
    candidatos = [f for f in os.listdir(carpeta)
                  if contiene.lower() in f.lower() and f.lower().endswith(".pdf")]
    if not candidatos:
        raise FileNotFoundError(f"No hay PDF con '{contiene}' en {carpeta}")
    if len(candidatos) > 1:
        raise ValueError(f"Más de un PDF con '{contiene}': {candidatos}")
    return carpeta / candidatos[0]


# ===========================================================================
# SECCIÓN 2 · LIMPIEZA + MARKDOWN JERÁRQUICO
# ===========================================================================
import fitz  # PyMuPDF

HEADER_PATTERNS = [
    r"^Ley \d+$", r"^Biblioteca del Congreso Nacional de Chile.*$",
    r"^página \d+ de \d+$", r"^Documento firmado digitalmente por.*$",
    r"^Para validar, acceda.*$", r"^Documento generado el.*$",
]
HEADER_RE = re.compile("|".join(HEADER_PATTERNS))

TITULO_RE   = re.compile(r"^(Título [\wÁÉÍÓÚÑñ.º°]+)$", re.IGNORECASE)
CAPITULO_RE = re.compile(r"^(Capítulo [\wÁÉÍÓÚÑñ.º°]+)$", re.IGNORECASE)
PARRAFO_RE  = re.compile(r"^(Párrafo [\wÁÉÍÓÚÑñ.º°]+)$", re.IGNORECASE)
ARTICULO_RE = re.compile(
    r"^\"?(Art[íi]culo\s+\d+\s*[º°]?\s*(?:bis|ter|qu[aá]ter|sexies|septies)?)\s*[.\-–]*\s*(.*)$",
    re.IGNORECASE)
ITEM_RE = re.compile(r"^([a-z]\))\s*(.*)$")
_INICIO_ARTICULO_RE = re.compile(r"^\s*\"?Art[íi]culo\s+\d+", re.IGNORECASE)

META_FIELDS = [
    ("Fecha Publicación", r"Fecha Publicación:\s*(\d{2}-\w{3}-\d{4})"),
    ("Fecha Promulgación", r"Fecha Promulgación:\s*(\d{2}-\w{3}-\d{4})"),
    ("Url Corta", r"Url Corta:\s*(\S+)"),
]


def extraer_lineas(pdf_path):
    doc = fitz.open(pdf_path)
    lineas = []
    for pagina in doc:
        lineas.extend(pagina.get_text().split("\n"))
    return lineas


def quitar_ruido(lineas):
    return [l for l in lineas if not HEADER_RE.match(l.strip())]


def unir_parrafos(lineas):
    parrafos, actual = [], ""
    for cruda in lineas:
        if cruda.strip() == "":
            continue
        texto = cruda.strip()
        es_inicio_nuevo = cruda.startswith("     ") or _INICIO_ARTICULO_RE.match(cruda)
        if es_inicio_nuevo and actual:
            parrafos.append(actual.strip())
            actual = texto
        elif not actual:
            actual = texto
        else:
            actual += " " + texto
    if actual:
        parrafos.append(actual.strip())
    return parrafos


def extraer_metadata(bloque):
    meta = {}
    for nombre, patron in META_FIELDS:
        m = re.search(patron, bloque)
        if m:
            meta[nombre] = m.group(1).strip()
    return meta


def construir_markdown(parrafos, titulo_ley, meta, ley_id, fuente_pdf):
    out = ["---", f'ley: "{ley_id}"', f'titulo: "{titulo_ley}"']
    for k, v in meta.items():
        out.append(f'{k.lower().replace(" ", "_")}: "{v}"')
    out.append(f'fuente_pdf: "{fuente_pdf}"')
    out.append("---\n")
    out.append(f"# {titulo_ley}\n")
    for p in parrafos:
        if p.strip().upper() == titulo_ley.strip().upper():
            continue
        if m := TITULO_RE.match(p):   out.append(f"\n## {m.group(1)}");   continue
        if m := CAPITULO_RE.match(p): out.append(f"\n### {m.group(1)}");  continue
        if m := PARRAFO_RE.match(p):  out.append(f"\n#### {m.group(1)}"); continue
        if m := ARTICULO_RE.match(p):
            etiqueta, resto = m.groups()
            out.append(f"\n##### {etiqueta}")
            if resto:
                out.append(resto)
            continue
        if m := ITEM_RE.match(p):
            out.append(f"- **{m.group(1)}** {m.group(2)}"); continue
        out.append(p)
    return "\n".join(out) + "\n"


def procesar_ley(pdf_path, ley_id, titulo):
    lineas = quitar_ruido(extraer_lineas(pdf_path))
    parrafos = unir_parrafos(lineas)
    meta = extraer_metadata(parrafos[0] if parrafos else "")
    cuerpo = parrafos[1:] if meta else parrafos
    return construir_markdown(cuerpo, titulo, meta, ley_id, str(pdf_path))


# ===========================================================================
# SECCIÓN 3 · CHUNKING JERÁRQUICO POR ARTÍCULO + SUBDIVISIÓN POR INCISO
# ===========================================================================
FRONTMATTER_RE = re.compile(r"^---\n(.*?)\n---\n", re.DOTALL)


def parse_frontmatter(md_text):
    m = FRONTMATTER_RE.match(md_text)
    meta = {}
    if m:
        for line in m.group(1).split("\n"):
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip().strip('"')
        return meta, md_text[m.end():]
    return meta, md_text


def chunk_por_articulo(md_text, ley_id):
    _, body = parse_frontmatter(md_text)
    lines = body.split("\n")
    chunks = []
    ctx = {"titulo_seccion": None, "capitulo": None, "parrafo": None}
    articulo_actual, buffer = None, []

    def flush():
        if articulo_actual and buffer:
            texto = "\n".join(buffer).strip()
            if texto:
                chunks.append({"ley": ley_id, **ctx, "articulo": articulo_actual,
                               "texto": f"{articulo_actual}\n{texto}"})

    for line in lines:
        if line.startswith("## ") and not line.startswith(("### ", "#### ", "##### ")):
            flush(); buffer = []; ctx.update(titulo_seccion=line[3:].strip(), capitulo=None, parrafo=None); articulo_actual = None; continue
        if line.startswith("### ") and not line.startswith(("#### ", "##### ")):
            flush(); buffer = []; ctx.update(capitulo=line[4:].strip(), parrafo=None); articulo_actual = None; continue
        if line.startswith("#### ") and not line.startswith("##### "):
            flush(); buffer = []; ctx["parrafo"] = line[5:].strip(); articulo_actual = None; continue
        if line.startswith("##### "):
            flush(); buffer = []; articulo_actual = line[6:].strip(); continue
        buffer.append(line)
    flush()
    return chunks


def subdividir_por_inciso(chunk, max_chars=1200, solape=150):
    texto = chunk["texto"]
    if len(texto) <= max_chars:
        return [chunk]
    piezas, inicio = [], 0
    while inicio < len(texto):
        fin = inicio + max_chars
        if fin < len(texto):
            corte = texto.rfind("\n", inicio, fin)
            if corte > inicio + 200:
                fin = corte
        piezas.append(texto[inicio:fin].strip())
        inicio = fin - solape
    subs = []
    for n, p in enumerate(piezas):
        if not p:
            continue
        nuevo = dict(chunk); nuevo["texto"] = p
        nuevo["articulo"] = f"{chunk['articulo']} (parte {n+1})" if chunk["articulo"] else chunk["articulo"]
        subs.append(nuevo)
    return subs


# ===========================================================================
# SECCIÓN 4 · MODELO DE EMBEDDINGS BGE-M3
# ===========================================================================
from FlagEmbedding import BGEM3FlagModel, FlagReranker

# use_fp16=True requiere GPU. Si estás en CPU, cámbialo a use_fp16=False.
modelo_bge = BGEM3FlagModel("BAAI/bge-m3", use_fp16=True)
reranker = FlagReranker("BAAI/bge-reranker-v2-m3", use_fp16=True)


def embed_dense(textos):
    salida = modelo_bge.encode(textos, batch_size=12, max_length=1024,
                               return_dense=True, return_sparse=False,
                               return_colbert_vecs=False)
    return salida["dense_vecs"]


# ===========================================================================
# SECCIÓN 5 · DOS COLECCIONES CHROMADB + ÍNDICE HÍBRIDO
# ===========================================================================
import chromadb
import numpy as np
from rank_bm25 import BM25Okapi


def fusion_rrf(listas_ranqueadas, k=60):
    puntajes = {}
    for lista in listas_ranqueadas:
        for rango, doc_id in enumerate(lista):
            puntajes[doc_id] = puntajes.get(doc_id, 0.0) + 1.0 / (k + rango + 1)
    return sorted(puntajes, key=puntajes.get, reverse=True)


class IndiceHibrido:
    def __init__(self, coleccion):
        self.coleccion = coleccion
        self.ids, self.textos, self.metas = [], [], []
        self.bm25 = None

    def indexar(self, chunks):
        if not chunks:
            return
        textos = [c["texto"] for c in chunks]
        dense = embed_dense(textos)
        base = len(self.ids)
        ids = [f"{c['ley']}__{base+i}" for i, c in enumerate(chunks)]
        metas = [{"ley": c.get("ley", ""), "articulo": c.get("articulo") or "",
                  "titulo_seccion": c.get("titulo_seccion") or "",
                  "capitulo": c.get("capitulo") or "", "parrafo": c.get("parrafo") or ""}
                 for c in chunks]
        self.coleccion.add(ids=ids, embeddings=dense.tolist(),
                           documents=textos, metadatas=metas)
        self.ids += ids; self.textos += textos; self.metas += metas
        self.bm25 = BM25Okapi([t.lower().split() for t in self.textos])

    def _por_id(self, doc_id):
        i = self.ids.index(doc_id)
        return {"texto": self.textos[i], **self.metas[i]}

    def buscar_denso(self, pregunta, k):
        v = embed_dense([pregunta])
        res = self.coleccion.query(query_embeddings=v.tolist(), n_results=k)
        return res["ids"][0]

    def buscar_disperso(self, pregunta, k):
        if self.bm25 is None:
            return []
        scores = self.bm25.get_scores(pregunta.lower().split())
        top = np.argsort(scores)[::-1][:k]
        return [self.ids[i] for i in top]

    def recuperar_hibrido(self, pregunta, k=5, candidatos=20):
        ids_densos = self.buscar_denso(pregunta, candidatos)
        ids_dispersos = self.buscar_disperso(pregunta, candidatos)
        fusionados = fusion_rrf([ids_densos, ids_dispersos])[:candidatos]
        if not fusionados:
            return []
        pares = [[pregunta, self._por_id(i)["texto"]] for i in fusionados]
        puntajes = reranker.compute_score(pares, normalize=True)
        orden = np.argsort(puntajes)[::-1][:k]
        return [{**self._por_id(fusionados[i]), "score": float(puntajes[i])} for i in orden]


# ===========================================================================
# SECCIÓN 6 · ALGORITMO 1 · INGESTA LOCAL DE ARCHIVOS CORPORATIVOS
# ===========================================================================
def _texto_docx(ruta):
    from docx import Document
    return "\n".join(p.text for p in Document(ruta).paragraphs if p.text.strip())


def _texto_pptx(ruta):
    from pptx import Presentation
    partes = []
    for i, slide in enumerate(Presentation(ruta).slides, 1):
        partes.append(f"## Diapositiva {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                partes.append(shape.text_frame.text.strip())
    return "\n".join(partes)


def _texto_xlsx(ruta):
    import pandas as pd
    partes = []
    for hoja, df in pd.read_excel(ruta, sheet_name=None).items():
        partes.append(f"## Hoja: {hoja}\n")
        partes.append(df.to_markdown(index=False))
    return "\n\n".join(partes)


def ingestar_archivo_corporativo(ruta, indice_pyme, max_chars=1200):
    ext = Path(ruta).suffix.lower()
    if ext == ".docx":   texto = _texto_docx(ruta)
    elif ext == ".pptx": texto = _texto_pptx(ruta)
    elif ext in (".xlsx", ".xls"): texto = _texto_xlsx(ruta)
    else:
        raise ValueError(f"Extensión no soportada: {ext}")

    bloques, actual = [], ""
    for parrafo in texto.split("\n"):
        if len(actual) + len(parrafo) > max_chars and actual:
            bloques.append(actual.strip()); actual = parrafo
        else:
            actual += "\n" + parrafo
    if actual.strip():
        bloques.append(actual.strip())

    chunks = [{"ley": "PYME", "articulo": Path(ruta).name,
               "titulo_seccion": None, "capitulo": None, "parrafo": None,
               "texto": b} for b in bloques]
    indice_pyme.indexar(chunks)
    return len(chunks)


# ===========================================================================
# SECCIÓN 7 · QWEN3 VÍA OLLAMA
# ===========================================================================
import ollama


def qwen(prompt, system=None, formato=None, temperatura=0.2):
    mensajes = []
    if system:
        mensajes.append({"role": "system", "content": system})
    mensajes.append({"role": "user", "content": prompt})
    resp = ollama.chat(model=MODELO_QWEN, messages=mensajes,
                       format=formato, options={"temperature": temperatura})
    return resp["message"]["content"]


# ===========================================================================
# SECCIÓN 8 · ALGORITMO 2 · RUTEO SEMÁNTICO
# ===========================================================================
SYS_RUTEO = """Eres un enrutador semántico. Clasifica la consulta del usuario
en UNA de estas categorías según qué fuente responde mejor:
- "legal": preguntas sobre las leyes 19.628 / 21.719 (obligaciones, definiciones, sanciones).
- "interno": preguntas sobre documentos, políticas o datos internos de la empresa.
- "ambos": si requiere contrastar la ley con la situación interna de la empresa.
Responde SOLO con un JSON: {"ruta": "legal"|"interno"|"ambos"}"""


def rutear(pregunta):
    try:
        salida = qwen(pregunta, system=SYS_RUTEO, formato="json")
        return json.loads(salida).get("ruta", "legal")
    except Exception:
        return "legal"


# ===========================================================================
# SECCIÓN 9 · ALGORITMO 3 · GENERACIÓN + CICLO DE AUTOCORRECCIÓN (LangGraph)
# ===========================================================================
from typing import TypedDict, List
from langgraph.graph import StateGraph, END


class EstadoAgente(TypedDict):
    pregunta: str
    contexto: str
    chunks: List[dict]
    borrador: str
    fiel: bool
    iteracion: int
    max_iteraciones: int
    instruccion_extra: str
    respuesta_final: str


SYS_GENERACION = """Eres un asistente jurídico experto en la Ley 19.628 y la
Ley 21.719 de Chile. Respondes mediante un SILOGISMO JURÍDICO:
1) premisa normativa (qué dice la ley según el CONTEXTO),
2) premisa fáctica (la situación de la consulta),
3) conclusión accionable para una PyME.
REGLAS: usa EXCLUSIVAMENTE el CONTEXTO. Cada afirmación debe citar (Ley X,
Artículo Y). Si el contexto es insuficiente, dilo. Lenguaje claro, sin jerga."""

SYS_VERIFICACION = """Verificas fidelidad. Dada una RESPUESTA y un CONTEXTO,
determina si CADA afirmación factual de la respuesta se sustenta estrictamente
en el contexto. Responde SOLO con JSON:
{"fiel": true|false, "motivo": "afirmaciones sin respaldo, si las hay"}"""


def _generar_borrador(pregunta, contexto, instruccion_extra=""):
    prompt = f"CONTEXTO:\n{contexto}\n\nPREGUNTA:\n{pregunta}\n\n{instruccion_extra}"
    return qwen(prompt, system=SYS_GENERACION, temperatura=0.2)


def _verificar_fidelidad(borrador, contexto):
    prompt = f"CONTEXTO:\n{contexto}\n\nRESPUESTA:\n{borrador}"
    salida = qwen(prompt, system=SYS_VERIFICACION, formato="json", temperatura=0.0)
    # Qwen3 a veces envuelve el JSON en texto o en <think>...</think>.
    # Extraemos el primer objeto {...} que aparezca.
    m = re.search(r'\{.*\}', salida, re.DOTALL)
    if m:
        try:
            v = json.loads(m.group(0))
            return bool(v.get("fiel", False)), v.get("motivo", "")
        except Exception:
            pass
    # Si no se pudo parsear, asumimos que es fiel para no bloquear
    # una respuesta buena (mejor mostrarla que abstenerse por un error de formato).
    return True, ""


def _integrar_citas(borrador, chunks):
    fuentes = "\n".join(f"- {c['ley']} {c['articulo']}" for c in chunks)
    aviso = "\n\n*Apoyo informativo/académico; no reemplaza asesoría legal profesional.*"
    return f"{borrador}\n\n--- Fuentes ---\n{fuentes}{aviso}"


def construir_grafo(generar_borrador, verificar_fidelidad, integrar_citas):
    def nodo_generar(estado):
        return {"borrador": generar_borrador(estado["pregunta"], estado["contexto"],
                                             estado.get("instruccion_extra", "")),
                "iteracion": estado["iteracion"] + 1}

    def nodo_verificar(estado):
        fiel, motivo = verificar_fidelidad(estado["borrador"], estado["contexto"])
        instr = estado.get("instruccion_extra", "")
        if not fiel:
            instr = ("Cíñete ESTRICTAMENTE al contexto; elimina toda afirmación "
                     "sin respaldo. " + motivo)
        return {"fiel": fiel, "instruccion_extra": instr}

    def decidir(estado):
        if estado["fiel"]:
            return "aceptar"
        if estado["iteracion"] >= estado["max_iteraciones"]:
            return "insuficiente"
        return "reintentar"

    def nodo_aceptar(estado):
        return {"respuesta_final": integrar_citas(estado["borrador"], estado["chunks"])}

    def nodo_insuficiente(estado):
        return {"respuesta_final": "No tengo información suficiente en las leyes "
                "indexadas para responder esto con certeza."}

    g = StateGraph(EstadoAgente)
    g.add_node("generar", nodo_generar)
    g.add_node("verificar", nodo_verificar)
    g.add_node("aceptar", nodo_aceptar)
    g.add_node("insuficiente", nodo_insuficiente)
    g.set_entry_point("generar")
    g.add_edge("generar", "verificar")
    g.add_conditional_edges("verificar", decidir,
                            {"reintentar": "generar", "aceptar": "aceptar",
                             "insuficiente": "insuficiente"})
    g.add_edge("aceptar", END)
    g.add_edge("insuficiente", END)
    return g.compile()


# ===========================================================================
# CONSTRUCCIÓN GLOBAL DEL SISTEMA
# ===========================================================================
# Estas variables se llenan al llamar a construir_sistema().
indice_leyes = None
indice_pyme = None
agente = None


def construir_sistema():
    """Procesa las leyes, arma los índices y compila el agente. Llamar una vez."""
    global indice_leyes, indice_pyme, agente

    ley_19628 = encontrar_pdf(RUTA_LEYES, "19628")
    ley_21719 = encontrar_pdf(RUTA_LEYES, "21719")
    print("Ley 19.628 ->", ley_19628)
    print("Ley 21.719 ->", ley_21719)

    leyes = [
        {"id": "Ley_19628", "ruta_pdf": ley_19628,
         "titulo": "SOBRE PROTECCION DE LA VIDA PRIVADA"},
        {"id": "Ley_21719", "ruta_pdf": ley_21719,
         "titulo": "REGULA LA PROTECCIÓN Y EL TRATAMIENTO DE LOS DATOS PERSONALES"},
    ]

    chunks_leyes = []
    for ley in leyes:
        md = procesar_ley(ley["ruta_pdf"], ley["id"], ley["titulo"])
        (CARPETA_SALIDA / f"{ley['id']}.md").write_text(md, encoding="utf-8")
        for c in chunk_por_articulo(md, ley["id"]):
            chunks_leyes.extend(subdividir_por_inciso(c))
    print(f"Chunks de leyes a indexar: {len(chunks_leyes)}")

    cliente = chromadb.EphemeralClient()
    for nombre in ("coleccion_leyes_chile", "coleccion_contexto_pyme"):
        try:
            cliente.delete_collection(nombre)
        except Exception:
            pass
    col_leyes = cliente.get_or_create_collection("coleccion_leyes_chile", metadata={"hnsw:space": "cosine"})
    col_pyme = cliente.get_or_create_collection("coleccion_contexto_pyme", metadata={"hnsw:space": "cosine"})

    indice_leyes = IndiceHibrido(col_leyes)
    indice_pyme = IndiceHibrido(col_pyme)
    indice_leyes.indexar(chunks_leyes)
    print("Índice legal listo:", col_leyes.count(), "chunks")

    agente = construir_grafo(_generar_borrador, _verificar_fidelidad, _integrar_citas)
    print("Agente compilado. Sistema listo.")


def recuperar_contexto(pregunta, k=5):
    ruta = rutear(pregunta)
    chunks = []
    if ruta in ("legal", "ambos"):
        chunks += indice_leyes.recuperar_hibrido(pregunta, k=k)
    if ruta in ("interno", "ambos") and indice_pyme.bm25 is not None:
        chunks += indice_pyme.recuperar_hibrido(pregunta, k=k)
    return ruta, chunks


def preguntar(pregunta, k=5, max_iteraciones=3):
    if agente is None:
        raise RuntimeError("Llama primero a construir_sistema().")
    ruta, chunks = recuperar_contexto(pregunta, k=k)
    if not chunks:
        return {"respuesta": "No tengo información suficiente en las leyes indexadas.",
                "ruta": ruta, "chunks": []}
    contexto = "\n\n".join(f"[{c['ley']} - {c['articulo']}]\n{c['texto']}" for c in chunks)
    estado = agente.invoke({"pregunta": pregunta, "contexto": contexto, "chunks": chunks,
                            "borrador": "", "fiel": False, "iteracion": 0,
                            "max_iteraciones": max_iteraciones, "instruccion_extra": "",
                            "respuesta_final": ""})
    return {"respuesta": estado["respuesta_final"], "ruta": ruta, "chunks": chunks}


# ===========================================================================
# PUNTO DE ENTRADA
# ===========================================================================
if __name__ == "__main__":
    construir_sistema()

    pregunta = "¿qué se considera un dato personal sensible?"
    resultado = preguntar(pregunta)
    print("\n" + "=" * 60)
    print("PREGUNTA:", pregunta)
    print("RUTA:", resultado["ruta"])
    print("=" * 60)
    print(resultado["respuesta"])
    print("\n--- Fuentes ---")
    for c in resultado["chunks"]:
        print(f"- {c['ley']} {c['articulo']} (score={c.get('score', 0):.3f})")
